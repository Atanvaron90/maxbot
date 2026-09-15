"""
MAX Bot — бот с поддержкой, информацией, условиями, рассылкой и проверкой долга.
Двухуровневая система обновления сверки из Google Таблиц:
  1) Триггеры (лист «Выгрузка продаж 5-ка») — ТОЛЬКО понедельник, раз в 15 мин,
     пока не обновятся в ОБОИХ таблицах; затем до следующего понедельника не проверяются.
     Зачёт = подмена сверки + рассылка всем пользователям.
  2) Лист «Сверка» — пн: раз в 15 мин, вт–пт: раз в 30 мин, сб–вс: не проверяется.
     Изменения подхватываются тихо (без уведомлений).
Автовосстановление: если кэш-файлы сверки отсутствуют/не читаются,
при запросе долга бот скачивает таблицы автоматически и повторяет поиск.
"""

import os
import sys
import json
import hashlib
import asyncio
import logging
import re
import math
import time
import urllib.request
import urllib.error
from datetime import datetime, timedelta, timezone
from enum import Enum
from typing import Dict, Optional, List
from dotenv import load_dotenv
from maxapi import Bot, Dispatcher, F
from maxapi.types import (
    MessageCreated, BotStarted, MessageCallback, Command,
    CallbackButton, ButtonsPayload, Attachment, BotCommand,
    InputMediaBuffer
)
from maxapi.enums.intent import Intent

try:
    import openpyxl
except ImportError:
    openpyxl = None
    logging.warning("openpyxl не установлен. Проверка долга не будет работать.")

# Московское время для всей логики сверки (на сервере локаль может быть UTC).
try:
    from zoneinfo import ZoneInfo
    MSK = ZoneInfo("Europe/Moscow")
except Exception:
    MSK = timezone(timedelta(hours=3), name="MSK")

load_dotenv()

DEBUG = os.getenv("DEBUG", "false").lower() == "true"

# FIX ENC: консоль Windows в cp1251 не умеет эмодзи — переводим потоки в UTF-8
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

logging.basicConfig(
    level=logging.DEBUG if DEBUG else logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("bot.log", encoding="utf-8"),  # FIX ENC: файл в UTF-8
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

BOT_TOKEN = os.getenv("MAX_BOT_TOKEN")
ADMIN_ID = int(os.getenv("ADMIN_ID", "0"))
USERS_FILE = 'data/users.json'

# ID Google Таблиц через запятую (одна, две или больше)
DEBT_SHEET_IDS: List[str] = [s.strip() for s in os.getenv("DEBT_SHEET_IDS", "").split(",") if s.strip()]
DEBT_DIR = os.getenv("DEBT_DIR", "data/debt")

# Имена листов
DEBT_DATA_SHEET = os.getenv("DEBT_DATA_SHEET", "Сверка")                    # лист с долгами
DEBT_SOURCE_SHEET = os.getenv("DEBT_SOURCE_SHEET", "Выгрузка продаж 5-ка")  # лист-триггер

# Интервалы проверки (минуты): понедельник — 15, вт–пт — 30, сб–вс — не проверяем
DEBT_CHECK_MONDAY_MINUTES = int(os.getenv("DEBT_CHECK_MONDAY_MINUTES", "15"))
DEBT_CHECK_WEEKDAY_MINUTES = int(os.getenv("DEBT_CHECK_WEEKDAY_MINUTES", "30"))

# Файл состояния (переживает перезапуски бота)
DEBT_STATE_FILE = os.getenv("DEBT_STATE_FILE", "data/debt_state.json")

REGION1_APP = os.getenv("REGION1_APP", "https://example.com/app1.jpg")
REGION1_QR = os.getenv("REGION1_QR", "https://example.com/qr1.jpg")
REGION1_BANK = os.getenv("REGION1_BANK", "https://example.com/bank1.jpg")
REGION2_APP = os.getenv("REGION2_APP", "https://example.com/app2.jpg")
REGION2_QR = os.getenv("REGION2_QR", "https://example.com/qr2.jpg")
REGION2_BANK = os.getenv("REGION2_BANK", "https://example.com/bank2.jpg")

if not BOT_TOKEN:
    logger.error("MAX_BOT_TOKEN не найден в .env")
    exit(1)

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

KNOWN_COMMANDS = {'/start', '/help', '/debt', '/change_name', '/cancel',
                  '/broadcast', '/reply', '/userinfo',
                  '/refresh_debt', '/debt_status'}


class UserState(Enum):
    IDLE = "idle"
    SUPPORT_MODE = "support"
    BROADCAST_MODE = "broadcast"
    DEBT_NAME_MODE = "debt_name"
    DEBT_CLARIFY_MODE = "debt_clarify"
    WAITING_FOR_FORWARD = "waiting_for_forward"
    ADMIN_DEBT_MODE = "admin_debt"
    ADMIN_DEBT_CLARIFY_MODE = "admin_debt_clarify"


user_states: Dict[int, UserState] = {}
user_clarify_data: Dict[int, Dict[str, object]] = {}
admin_debt_data: Dict[int, Dict[str, object]] = {}

_debt_cache: Dict = {'mtime': None, 'rows': None, 'loaded_at': None}

_debt_state: Dict = {
    'trig_baseline': None,      # сигнатуры листа-триггера на момент последнего зачёта
    'data_baseline': None,      # сигнатуры листа «Сверка», соответствующие активному кэшу
    'trig_committed_on': None,  # 'YYYY-MM-DD' — когда триггеры были зачтены (пн-блокировка)
    'last_sigs_t': None,
    'last_sigs_d': None,
    'last_check': None,
}


# ---------- РАБОТА С БАЗОЙ ПОЛЬЗОВАТЕЛЕЙ ----------

def load_users() -> Dict:
    try:
        os.makedirs('data', exist_ok=True)
        if os.path.exists(USERS_FILE):
            with open(USERS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception as e:
        logger.error(f"Ошибка загрузки: {e}")
    return {}


def save_users(users: Dict):
    try:
        os.makedirs('data', exist_ok=True)
        with open(USERS_FILE, 'w', encoding='utf-8') as f:
            json.dump(users, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.error(f"Ошибка сохранения: {e}")


def add_user(user_id: int, chat_id: int, username: str = None, name: str = None, full_name: str = None):
    users = load_users()
    uid = str(user_id)
    is_new = uid not in users
    now_msk = datetime.now(MSK).isoformat()

    if is_new:
        users[uid] = {
            'chat_id': chat_id,
            'username': username,
            'name': name,
            'full_name': full_name,
            'created_at': now_msk,
            'last_activity': now_msk
        }
    else:
        existing = users[uid]
        if chat_id is not None:
            existing['chat_id'] = chat_id
        if username is not None:
            existing['username'] = username
        if name is not None:
            existing['name'] = name
        if full_name is not None:
            existing['full_name'] = full_name
        existing['last_activity'] = now_msk
        if 'created_at' not in existing:
            existing['created_at'] = now_msk
        users[uid] = existing

    save_users(users)
    logger.info(f"Обновлён пользователь: id={user_id}, is_new={is_new}")
    return is_new


def ensure_user(user_id: int, chat_id: int, username: str = None, name: str = None):
    return add_user(user_id, chat_id, username, name)


def set_user_full_name(user_id: int, full_name: str):
    users = load_users()
    uid = str(user_id)
    if uid in users:
        users[uid]['full_name'] = full_name
        save_users(users)
        logger.info(f"Сохранено ФИО для user_id={user_id}")
        return True
    logger.warning(f"Не удалось сохранить ФИО для user_id={user_id}: пользователь не найден")
    return False


def clear_user_full_name(user_id: int):
    users = load_users()
    uid = str(user_id)
    if uid in users:
        users[uid]['full_name'] = None
        save_users(users)
        logger.info(f"Очищено ФИО для user_id={user_id}")
        return True
    return False


# ---------- МЕНЮ ----------

def create_main_menu(user_id: int) -> Attachment:
    row1 = [
        CallbackButton(text="📋 Условия", payload="conditions", intent=Intent.DEFAULT),
        CallbackButton(text="💳 Долг", payload="debt", intent=Intent.DEFAULT)
    ]
    row2 = [
        CallbackButton(text="ℹ️ О боте", payload="info", intent=Intent.DEFAULT),
        CallbackButton(text="💬 Поддержка", payload="support", intent=Intent.POSITIVE)
    ]
    row3 = [CallbackButton(text="❓ Справка", payload="help_command", intent=Intent.DEFAULT)]
    if user_id == ADMIN_ID:
        row3.insert(0, CallbackButton(text="📢 Новости", payload="broadcast", intent=Intent.DEFAULT))

    rows = [row1, row2, row3]

    if user_id == ADMIN_ID:
        rows.append([
            CallbackButton(text="👤 Данные контакта", payload="userinfo", intent=Intent.DEFAULT)
        ])

    buttons = ButtonsPayload(buttons=rows)
    return Attachment(type="inline_keyboard", payload=buttons)


def create_debt_result_menu() -> Attachment:
    buttons = ButtonsPayload(buttons=[
        [CallbackButton(text="🏠 Главное меню", payload="back_to_main", intent=Intent.DEFAULT)],
        [CallbackButton(text="🔄 Сменить ФИО", payload="change_name", intent=Intent.DEFAULT)]
    ])
    return Attachment(type="inline_keyboard", payload=buttons)


def create_admin_debt_menu() -> Attachment:
    buttons = ButtonsPayload(buttons=[
        [CallbackButton(text="🔍 Проверить другое ФИО", payload="admin_debt_again", intent=Intent.DEFAULT)],
        [CallbackButton(text="🏠 Главное меню", payload="back_to_main", intent=Intent.DEFAULT)],
    ])
    return Attachment(type="inline_keyboard", payload=buttons)


def create_regions_menu() -> Attachment:
    buttons = ButtonsPayload(buttons=[
        [
            CallbackButton(text="📍 Краснодар", payload="region_1", intent=Intent.DEFAULT),
            CallbackButton(text="📍 КМВ, КЧР", payload="region_2", intent=Intent.DEFAULT)
        ],
        [
            CallbackButton(text="🔙 Назад", payload="back_to_main", intent=Intent.DEFAULT)
        ]
    ])
    return Attachment(type="inline_keyboard", payload=buttons)


def create_materials_menu(region: str) -> Attachment:
    buttons = ButtonsPayload(buttons=[
        [CallbackButton(text="📄 Заявление", payload=f"{region}_app", intent=Intent.DEFAULT)],
        [CallbackButton(text="📱 QR-код", payload=f"{region}_qr", intent=Intent.DEFAULT)],
        [CallbackButton(text="💳 Реквизиты", payload=f"{region}_bank", intent=Intent.DEFAULT)],
        [CallbackButton(text="🔙 Назад к регионам", payload="back_to_regions", intent=Intent.DEFAULT)]
    ])
    return Attachment(type="inline_keyboard", payload=buttons)


def create_cancel_menu() -> Attachment:
    buttons = ButtonsPayload(buttons=[
        [CallbackButton(text="❌ Отмена", payload="cancel_action", intent=Intent.DEFAULT)]
    ])
    return Attachment(type="inline_keyboard", payload=buttons)


# ---------- ФУНКЦИИ ДЛЯ РАБОТЫ С ДОЛГОМ ----------

def parse_balance(value) -> float:
    """Надёжный парсинг баланса: числа, '1 234,56 р.', '-500' и т.д."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    s = re.sub(r'[^\d,.\-]', '', str(value))
    if not s.strip('-'):
        return 0.0

    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    else:
        s = s.replace(',', '.')
        if s.count('.') > 1:
            head, _, tail = s.rpartition('.')
            s = head.replace('.', '') + '.' + tail

    try:
        return float(s)
    except ValueError:
        logger.warning(f"Не удалось преобразовать значение в число: {value!r}")
        return 0.0


def _norm_word(w: str) -> str:
    return w.lower().replace('ё', 'е')


def clean_name(name: str) -> str:
    name = name.strip()
    if '(' in name:
        name = name.split('(')[0].strip()
    if ',' in name:
        name = name.split(',')[0].strip()
    return name


# FIX SHEET: нормализация имён листов — устойчивость к невидимым символам
def _norm_sheet_name(name: str) -> str:
    """Нормализация имени листа: неразрывные пробелы → обычные,
    схлопывание повторных пробелов, унификация тире, lower, strip.
    Позволяет находить лист даже при невидимых символах в названии."""
    s = str(name).replace('\xa0', ' ').replace('\u2007', ' ').replace('\u202f', ' ')
    s = re.sub(r'\s+', ' ', s).strip().lower()
    for dash in ('\u2010', '\u2011', '\u2012', '\u2013', '\u2014', '\u2212'):
        s = s.replace(dash, '-')
    return s


def _debt_paths() -> List[str]:
    return [os.path.join(DEBT_DIR, f"sheet_{i}.xlsx") for i in range(1, len(DEBT_SHEET_IDS) + 1)]


def _load_debt_rows() -> Optional[List[Dict]]:
    """Читает лист «Сверка» из каждого кэш-файла (другие листы игнорируются).
    None — техническая ошибка (файлы отсутствуют/не читаются), [] — данные пусты."""
    if not openpyxl:
        logger.error("openpyxl не установлен")
        return None

    paths = _debt_paths()
    if not paths:
        logger.error("Список таблиц пуст (DEBT_SHEET_IDS)")
        return None

    existing = [p for p in paths if os.path.exists(p)]
    if not existing:
        logger.error("Ни одного кэш-файла сверки не найдено")
        return None

    try:
        mtimes = tuple(os.path.getmtime(p) for p in existing)
    except OSError as e:
        logger.error(f"Не удалось получить mtime: {e}")
        return None

    if _debt_cache['mtime'] == mtimes and _debt_cache['rows'] is not None:
        return _debt_cache['rows']

    last_error = None
    for attempt in range(3):
        try:
            rows: List[Dict] = []
            target_norm = _norm_sheet_name(DEBT_DATA_SHEET)
            for path in existing:
                wb = openpyxl.load_workbook(path, data_only=True)
                # FIX SHEET: ищем лист «Сверка» с нормализацией имён
                sheet_name = next((s for s in wb.sheetnames if _norm_sheet_name(s) == target_norm), None)
                if sheet_name is None:
                    logger.warning(
                        f"В файле {os.path.basename(path)} нет листа «{DEBT_DATA_SHEET}». "
                        f"Реальные имена: {[repr(s) for s in wb.sheetnames]}"
                    )
                    wb.close()
                    continue
                sheet = wb[sheet_name]
                for row in sheet.iter_rows(min_row=2, max_col=5):
                    cell_name = row[0].value
                    if not cell_name:
                        continue
                    rows.append({
                        'name': str(cell_name).strip(),
                        'balance': parse_balance(row[4].value),
                        'sheet': sheet_name,
                        'file': os.path.basename(path),
                    })
                wb.close()

            _debt_cache['mtime'] = mtimes
            _debt_cache['rows'] = rows
            _debt_cache['loaded_at'] = datetime.now(MSK)
            logger.info(f"Лист «{DEBT_DATA_SHEET}» прочитан: {len(rows)} строк из {len(existing)} файл(ов)")
            return rows
        except Exception as e:
            last_error = e
            logger.warning(f"Попытка {attempt + 1}/3 прочитать файлы долгов не удалась: {e}")
            time.sleep(0.5)

    logger.error(f"Не удалось прочитать файлы долгов после 3 попыток: {last_error}")
    if _debt_cache['rows'] is not None:
        logger.warning("Использую предыдущую версию данных сверки")
        return _debt_cache['rows']
    return None


def search_debt_by_name(full_name: str) -> Optional[List[Dict]]:
    rows = _load_debt_rows()
    if rows is None:
        return None

    words = full_name.strip().split()
    if len(words) < 2:
        return []

    n = min(len(words), 3)
    search_words = [_norm_word(w) for w in words[:n]]
    matches = []

    for item in rows:
        cell_words = clean_name(item['name']).split()
        if len(cell_words) < n:
            continue
        cell_search = [_norm_word(w) for w in cell_words[:n]]
        if search_words == cell_search:
            matches.append(item)

    return matches


def group_matches(matches: List[Dict]) -> Dict[str, List[Dict]]:
    groups: Dict[str, List[Dict]] = {}
    for item in matches:
        clean = clean_name(item['name'])
        key = ' '.join(_norm_word(w) for w in clean.split()[:3])
        groups.setdefault(key, []).append(item)
    return groups


def get_last_sunday() -> str:
    # В понедельник до 13:00 МСК сверка ещё считается прошлой неделей,
    # после 13:00 — свежей (пн–вс).
    now = datetime.now(MSK)
    today = now.date()
    if today.weekday() == 0 and now.hour < 13:
        delta = timedelta(days=8)
    else:
        if today.weekday() == 6:
            delta = timedelta(days=7)
        else:
            days_to_subtract = today.weekday() + 1
            delta = timedelta(days=days_to_subtract)
    last_sunday = today - delta
    return last_sunday.strftime("%d.%m.%Y")


def format_debt_result(matches: List[Dict], date_str: str) -> str:
    groups = group_matches(matches)

    if len(groups) > 1:
        all_negative = [item for items in groups.values() for item in items if item['balance'] < 0]
        if all_negative:
            lines = [f"⚠️ **Найдены задолженности по разным магазинам (по {date_str} включительно):**"]
            for idx, item in enumerate(all_negative, 1):
                rounded = math.ceil(abs(item['balance']))
                lines.append(f"{idx}. {item['name']} — задолженность: {rounded} руб.")
            lines.append("\n_При необходимости можете уточнить информацию у менеджера._")
            return "\n".join(lines)
        return f"✅ **По {date_str} включительно у вас нет задолженности.**"

    group_items = next(iter(groups.values()))
    negative_items = [item for item in group_items if item['balance'] < 0]

    if not negative_items:
        return f"✅ **По {date_str} включительно у вас нет задолженности.**"
    elif len(negative_items) == 1:
        rounded = math.ceil(abs(negative_items[0]['balance']))
        return f"⚠️ **По {date_str} включительно у вас есть задолженность на сумму {rounded} руб.**"
    else:
        lines = [f"⚠️ **Найдены задолженности по разным магазинам (по {date_str} включительно):**"]
        for idx, item in enumerate(negative_items, 1):
            rounded = math.ceil(abs(item['balance']))
            lines.append(f"{idx}. {item['name']} — задолженность: {rounded} руб.")
        lines.append("\n_При необходимости можете уточнить информацию у менеджера._")
        return "\n".join(lines)


async def send_debt_result(chat_id: int, user_id: int, matches: List[Dict],
                           menu: Optional[Attachment] = None):
    date_str = get_last_sunday()
    result_text = format_debt_result(matches, date_str)
    await bot.send_message(
        chat_id=chat_id,
        text=result_text,
        format="markdown",
        attachments=[menu if menu is not None else create_debt_result_menu()]
    )


# ---------- СИСТЕМА ОБНОВЛЕНИЯ СВЕРКИ ----------

# Таймаут и число попыток скачивания (можно переопределить в .env)
DOWNLOAD_TIMEOUT = int(os.getenv("DOWNLOAD_TIMEOUT", "120"))   # секунд
DOWNLOAD_RETRIES = int(os.getenv("DOWNLOAD_RETRIES", "3"))


def _download_debt_file_sync(url: str, tmp_path: str) -> None:
    """Синхронное скачивание с ретраями (выполняется в отдельном потоке).
    Google генерирует xlsx для больших таблиц с паузой до минуты-другой,
    поэтому таймаут по умолчанию 120 с и до 3 попыток."""
    last_error: Exception = None
    for attempt in range(1, DOWNLOAD_RETRIES + 1):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=DOWNLOAD_TIMEOUT) as resp:
                data = resp.read()

            if not data.startswith(b"PK"):
                raise ValueError("Ответ не является xlsx-файлом (проверьте доступ к таблице)")

            with open(tmp_path, "wb") as f:
                f.write(data)
            return

        except urllib.error.HTTPError as e:
            # FIX HTTP: человекочитаемая подсказка; повторять бессмысленно — сразу наверх
            if e.code in (401, 403):
                raise ValueError(
                    f"HTTP {e.code}: таблица закрыта доступом. Откройте «Настройки доступа» → "
                    "«Все, у кого есть ссылка» → Читатель."
                ) from e
            if e.code == 404:
                raise ValueError("HTTP 404: таблица не найдена — проверьте DEBT_SHEET_IDS.") from e
            last_error = e
            logger.warning(f"Скачивание, попытка {attempt}/{DOWNLOAD_RETRIES}: HTTP {e.code}")
        except Exception as e:
            # сетевые сбои и таймауты — повторяем
            last_error = e
            logger.warning(f"Скачивание, попытка {attempt}/{DOWNLOAD_RETRIES} не удалась: {e}")

        if attempt < DOWNLOAD_RETRIES:
            time.sleep(5 * attempt)  # пауза 5с, затем 10с

    raise last_error


def _cleanup_files(paths: List[str]):
    for p in paths:
        try:
            if os.path.exists(p):
                os.remove(p)
        except OSError:
            pass


async def _download_tables_tmp() -> Optional[List[str]]:
    """Скачивает таблицы ПОСЛЕДОВАТЕЛЬНО во временные файлы.
    FIX: временный файл называется sheet_N.tmp.xlsx (а не sheet_N.xlsx.tmp) —
    openpyxl валидирует расширение и отказывается открывать '.tmp'.
    None — если хотя бы одна не скачалась (кэш при этом не трогаем)."""
    tmps = [p.replace(".xlsx", ".tmp.xlsx") for p in _debt_paths()]
    os.makedirs(DEBT_DIR, exist_ok=True)

    try:
        for sid, path in zip(DEBT_SHEET_IDS, tmps):
            url = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
            await asyncio.to_thread(_download_debt_file_sync, url, path)
            logger.info(f"Таблица {sid[:10]}… скачана")
        return tmps
    except Exception as e:
        logger.error(f"Не удалось скачать таблицу: {e}")
        _cleanup_files(tmps)
        return None


def _sheet_signature(path: str, sheet_name: str) -> Optional[str]:
    """md5-отпечаток содержимого конкретного листа. None — лист не найден/файл битый."""
    if not openpyxl:
        return None

    # FIX: если read_only-режим падает (бывает на Google-экспорте), пробуем
    # открыть файл обычным способом и обязательно логируем причину
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        logger.warning(f"openpyxl (read_only) не смог открыть {os.path.basename(path)}: {e}")
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e2:
            logger.warning(f"Не удалось открыть {os.path.basename(path)} совсем: {e2}")
            return None

    try:
        target = _norm_sheet_name(sheet_name)
        name = next((s for s in wb.sheetnames if _norm_sheet_name(s) == target), None)
        if name is None:
            logger.warning(
                f"Лист «{sheet_name}» не найден в {os.path.basename(path)}. "
                f"Реальные имена листов: {[repr(s) for s in wb.sheetnames]}"
            )
            return None
        h = hashlib.md5()
        for row in wb[name].iter_rows(values_only=True):
            h.update(repr(row).encode('utf-8'))
        return h.hexdigest()
    finally:
        wb.close()


def _signatures(tmps: List[str], sheet_name: str) -> List[Optional[str]]:
    return [_sheet_signature(t, sheet_name) for t in tmps]


def _commit_tables(tmps: List[str]):
    """Атомарно подменяет кэш-файлы сверки (парой)."""
    for tmp, path in zip(tmps, _debt_paths()):
        os.replace(tmp, path)
    logger.info("Кэш сверки обновлён")


# Персистентность состояния
def _load_debt_state():
    try:
        if os.path.exists(DEBT_STATE_FILE):
            with open(DEBT_STATE_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            n = len(DEBT_SHEET_IDS)
            tb = data.get('trig_baseline')
            db = data.get('data_baseline')
            if isinstance(tb, list) and len(tb) == n:
                _debt_state['trig_baseline'] = tb
            if isinstance(db, list) and len(db) == n:
                _debt_state['data_baseline'] = db
            tco = data.get('trig_committed_on')
            if isinstance(tco, str):
                _debt_state['trig_committed_on'] = tco
    except Exception as e:
        logger.warning(f"Не удалось прочитать состояние сверки: {e}")


def _save_debt_state():
    try:
        os.makedirs('data', exist_ok=True)
        with open(DEBT_STATE_FILE, 'w', encoding='utf-8') as f:
            json.dump({
                'trig_baseline': _debt_state['trig_baseline'],
                'data_baseline': _debt_state['data_baseline'],
                'trig_committed_on': _debt_state['trig_committed_on'],
            }, f, ensure_ascii=False)
    except Exception as e:
        logger.error(f"Не удалось сохранить состояние сверки: {e}")


async def notify_debt_updated():
    """Рассылает сообщение об обновлении сверки ВСЕМ пользователям из базы."""
    text = "Сверка обновлена. Можно проверить наличие актуальной задолженности."
    users = load_users()
    sent = 0
    errors = 0
    for uid, user_data in users.items():
        chat_id_user = user_data.get('chat_id')
        if not chat_id_user:
            continue
        try:
            await bot.send_message(chat_id=chat_id_user, text=text)
            sent += 1
        except Exception as e:
            errors += 1
            logger.error(f"Уведомление не доставлено пользователю {uid}: {e}")
        await asyncio.sleep(0.2)
    logger.info(f"Уведомление об обновлении сверки: отправлено {sent}, ошибок {errors}")


# ---------- АВТОВОССТАНОВЛЕНИЕ КЭША ПРИ ЗАПРОСЕ ДОЛГА ----------

_debt_recover_lock: Optional[asyncio.Lock] = None


async def ensure_debt_cache() -> bool:
    """Восстанавливает кэш сверки скачиванием, если файлы отсутствуют
    или не читаются. True — после вызова данные доступны.
    Блокировка защищает от параллельных скачиваний при одновременных запросах."""
    global _debt_recover_lock
    if _debt_recover_lock is None:
        _debt_recover_lock = asyncio.Lock()

    async with _debt_recover_lock:
        # Повторная проверка: пока ждали блокировку, другой запрос мог уже всё восстановить
        if _load_debt_rows() is not None:
            return True

        logger.warning("Кэш сверки недоступен — скачиваю таблицы по запросу")
        tmps = await _download_tables_tmp()
        if tmps is None:
            logger.error("Восстановить кэш сверки не удалось (таблицы не скачались)")
            return False

        trig_sigs = _signatures(tmps, DEBT_SOURCE_SHEET)
        data_sigs = _signatures(tmps, DEBT_DATA_SHEET)

        _commit_tables(tmps)

        # Синхронизируем базовые сигнатуры: активный кэш теперь равен текущему
        # состоянию Google. Если состояния ещё нет вовсе — фиксируем как базовое,
        # чтобы первый же плановый чек не сгенерировал лишнее уведомление.
        if all(s is not None for s in data_sigs):
            _debt_state['data_baseline'] = data_sigs[:]
            if _debt_state['trig_baseline'] is None and all(s is not None for s in trig_sigs):
                _debt_state['trig_baseline'] = trig_sigs[:]
        _save_debt_state()

        _load_debt_rows()  # прогреть кэш прочитанных строк
        logger.info("Кэш сверки восстановлен по запросу")
        return True


async def search_debt_with_recovery(full_name: str) -> Optional[List[Dict]]:
    """Поиск долга с автовосстановлением. Если кэш недоступен (файлы
    отсутствуют или не читаются) — скачивает таблицы и повторяет поиск.
    None возвращается только если восстановление тоже не помогло."""
    matches = search_debt_by_name(full_name)
    if matches is None:
        if await ensure_debt_cache():
            matches = search_debt_by_name(full_name)
    return matches


# ---------- ПЛАНОВЫЕ ПРОВЕРКИ ----------

async def startup_check() -> str:
    """Проверка при старте бота. Восстанавливает состояние из файла;
    если триггеры обновились, пока бот был выключен — зачитывает с рассылкой;
    если менялся только лист «Сверка» — тихо обновляет кэш."""
    _load_debt_state()
    n = len(DEBT_SHEET_IDS)

    tmps = await _download_tables_tmp()
    if tmps is None:
        return "❌ Не удалось скачать одну из таблиц (см. лог)."

    trig_sigs = _signatures(tmps, DEBT_SOURCE_SHEET)
    data_sigs = _signatures(tmps, DEBT_DATA_SHEET)
    if any(s is None for s in trig_sigs) or any(s is None for s in data_sigs):
        _cleanup_files(tmps)
        return (f"⚠️ Не найден лист «{DEBT_SOURCE_SHEET}» или «{DEBT_DATA_SHEET}». "
                "Точные имена листов см. в предупреждениях выше в логе. "
                "Проверьте DEBT_DATA_SHEET / DEBT_SOURCE_SHEET в настройках.")

    st = _debt_state
    st['last_sigs_t'] = trig_sigs
    st['last_sigs_d'] = data_sigs
    st['last_check'] = datetime.now(MSK)

    # Первый запуск вообще (нет состояния) — фиксируем как базовое, без уведомлений
    if st['trig_baseline'] is None or st['data_baseline'] is None:
        _commit_tables(tmps)
        st['trig_baseline'] = trig_sigs[:]
        st['data_baseline'] = data_sigs[:]
        _save_debt_state()
        rows = _load_debt_rows()
        count = len(rows) if rows is not None else "?"
        return f"✅ Сверка загружена ({count} записей). Базовое состояние зафиксировано."

    today = datetime.now(MSK).date().isoformat()

    # Триггеры обновились, пока бот был выключен → зачёт с рассылкой (один раз)
    if all(trig_sigs[i] != st['trig_baseline'][i] for i in range(n)):
        _commit_tables(tmps)
        st['trig_baseline'] = trig_sigs[:]
        st['data_baseline'] = data_sigs[:]
        st['trig_committed_on'] = today
        _save_debt_state()
        rows = _load_debt_rows()
        count = len(rows) if rows is not None else "?"
        await notify_debt_updated()
        return (f"✅ За время отключения таблицы обновились. Сверка актуальна "
                f"({count} записей), уведомление отправлено пользователям.")

    # Только сверка изменилась → тихая подмена
    if any(data_sigs[i] != st['data_baseline'][i] for i in range(n)):
        _commit_tables(tmps)
        st['data_baseline'] = data_sigs[:]
        _save_debt_state()
        rows = _load_debt_rows()
        count = len(rows) if rows is not None else "?"
        return f"🔄 Лист «{DEBT_DATA_SHEET}» обновился — сверка актуализирована ({count} записей)."

    _cleanup_files(tmps)
    rows = _load_debt_rows()
    count = len(rows) if rows is not None else "?"
    return f"✅ Сверка актуальна ({count} записей), изменений нет."


async def scheduled_check(force_weekday: bool = False) -> str:
    """Плановая проверка.
    1) Триггеры: только в понедельник (или force_weekday для ручного /refresh_debt)
       и только если ещё не зачтены сегодня. Оба изменились → зачёт + рассылка.
    2) Сверка: пн–пт; изменилась → тихая подмена кэша."""
    n = len(DEBT_SHEET_IDS)

    tmps = await _download_tables_tmp()
    if tmps is None:
        return "❌ Не удалось скачать одну из таблиц (см. лог). Пробую при следующей проверке."

    trig_sigs = _signatures(tmps, DEBT_SOURCE_SHEET)
    data_sigs = _signatures(tmps, DEBT_DATA_SHEET)
    if any(s is None for s in trig_sigs) or any(s is None for s in data_sigs):
        _cleanup_files(tmps)
        return (f"⚠️ Не найден лист «{DEBT_SOURCE_SHEET}» или «{DEBT_DATA_SHEET}». "
                "Точные имена листов см. в предупреждениях выше в логе. "
                "Проверьте DEBT_DATA_SHEET / DEBT_SOURCE_SHEET в настройках.")

    now = datetime.now(MSK)
    today = now.date().isoformat()
    wd = now.weekday()

    st = _debt_state
    st['last_sigs_t'] = trig_sigs
    st['last_sigs_d'] = data_sigs
    st['last_check'] = now

    lines: List[str] = []
    trig_done = False

    # ---------- 1. Триггерная проверка (пн, до зачёта) ----------
    if (wd == 0 or force_weekday) and st['trig_committed_on'] != today:
        if st['trig_baseline'] is None or len(st['trig_baseline']) != n:
            st['trig_baseline'] = trig_sigs[:]
        elif all(trig_sigs[i] != st['trig_baseline'][i] for i in range(n)):
            # ОБА триггера обновились — зачёт
            _commit_tables(tmps)
            st['trig_baseline'] = trig_sigs[:]
            st['data_baseline'] = data_sigs[:]
            st['trig_committed_on'] = today
            trig_done = True
            rows = _load_debt_rows()
            count = len(rows) if rows is not None else "?"
            logger.info(f"Недельное обновление сверки зачтено ({count} записей)")
            await notify_debt_updated()
            lines.append(f"✅ Сверка обновлена и разослана всем пользователям ({count} записей).")
        else:
            changed = [i + 1 for i in range(n) if trig_sigs[i] != st['trig_baseline'][i]]
            waiting = [i + 1 for i in range(n) if trig_sigs[i] == st['trig_baseline'][i]]
            if changed:
                lines.append(f"⏳ Триггер(ы) обновились в таблице {changed}; жду таблицу {waiting}. "
                             "Рассылка будет после обновления всех таблиц.")

    # ---------- 2. Сверка-проверка (пн–пт, тихо) ----------
    if not trig_done:
        if wd in (5, 6) and not force_weekday:
            _cleanup_files(tmps)  # в выходные плановых сверок не делаем
        elif st['data_baseline'] is None or len(st['data_baseline']) != n:
            _commit_tables(tmps)
            st['data_baseline'] = data_sigs[:]
            if st['trig_baseline'] is None:
                st['trig_baseline'] = trig_sigs[:]
            rows = _load_debt_rows()
            count = len(rows) if rows is not None else "?"
            lines.append(f"🔄 Лист «{DEBT_DATA_SHEET}» зафиксирован ({count} записей).")
        elif any(data_sigs[i] != st['data_baseline'][i] for i in range(n)):
            changed = [i + 1 for i in range(n) if data_sigs[i] != st['data_baseline'][i]]
            _commit_tables(tmps)
            st['data_baseline'] = data_sigs[:]
            rows = _load_debt_rows()
            count = len(rows) if rows is not None else "?"
            lines.append(f"🔄 Лист «{DEBT_DATA_SHEET}» изменился в таблице(ах) {changed} — "
                         f"активная сверка обновлена ({count} записей), без рассылки.")
        else:
            _cleanup_files(tmps)
            if not lines:
                lines.append("— Изменений в таблицах нет.")

    _save_debt_state()
    return "\n".join(lines) if lines else "— Изменений нет."


async def debt_refresh_loop():
    """Расписание фоновых проверок (московское время):
    • стартовая проверка — всегда при запуске;
    • пн: каждые DEBT_CHECK_MONDAY_MINUTES (триггеры до зачёта + сверка);
    • вт–пт: каждые DEBT_CHECK_WEEKDAY_MINUTES (только сверка);
    • сб–вс: плановых проверок нет."""
    if not DEBT_SHEET_IDS:
        logger.warning("DEBT_SHEET_IDS не задан — обновление сверки отключено")
        return

    try:
        status = await startup_check()
        logger.info(f"Стартовая загрузка сверки: {status}")
    except Exception as e:
        logger.error(f"Ошибка стартовой загрузки сверки: {e}")

    while True:
        now = datetime.now(MSK)
        wd = now.weekday()

        if wd in (5, 6):
            await asyncio.sleep(3600)
            continue

        interval_min = DEBT_CHECK_MONDAY_MINUTES if wd == 0 else DEBT_CHECK_WEEKDAY_MINUTES
        await asyncio.sleep(interval_min * 60)

        try:
            status = await scheduled_check()
            logger.info(f"Плановая проверка сверки: {status}")
        except Exception as e:
            logger.error(f"Ошибка плановой проверки сверки: {e}")


# ---------- ОБРАБОТКА USERINFO ----------

def format_user_info_from_dict(user_id: int, user_data: Dict) -> str:
    fw_chat_id = user_data.get('chat_id', '—')
    fw_name = user_data.get('name', '—')

    return (
        "👤 **Данные контакта:**\n\n"
        f"• **User ID:** `{user_id}`\n"
        f"• **Chat ID:** `{fw_chat_id}`\n"
        f"• **Имя:** {fw_name or '—'}\n"
    )


def format_forwarded_user_info(sender) -> str:
    fw_user_id = getattr(sender, 'user_id', None)
    fw_chat_id = getattr(sender, 'chat_id', None)
    fw_first_name = getattr(sender, 'first_name', None)
    fw_last_name = getattr(sender, 'last_name', None)

    if not fw_chat_id and fw_user_id:
        users = load_users()
        udata = users.get(str(fw_user_id))
        if udata:
            fw_chat_id = udata.get('chat_id')

    fio = " ".join(filter(None, [fw_last_name, fw_first_name])) or '—'

    return (
        "👤 **Данные контакта:**\n\n"
        f"• **User ID:** `{fw_user_id if fw_user_id else '—'}\n"
        f"• **Chat ID:** `{fw_chat_id if fw_chat_id else '—'}`\n"
        f"• **Имя:** {fio}\n"
    )


async def process_userinfo_input(event, user_id: int, chat_id: int, text: str):
    link = getattr(event.message, 'link', None)
    if link and getattr(link, 'sender', None):
        sender = link.sender
        logger.info(f"Пересылка обнаружена, отправитель: user_id={getattr(sender, 'user_id', None)}")
        text_out = format_forwarded_user_info(sender)
        await bot.send_message(
            chat_id=chat_id,
            text=text_out,
            format="markdown",
            attachments=[create_main_menu(user_id)]
        )
        user_states.pop(user_id, None)
        return

    users = load_users()
    user_data = users.get(str(user_id))
    if user_data:
        text_out = format_user_info_from_dict(user_id, user_data)
    else:
        text_out = (
            "👤 **Ваши данные:**\n\n"
            f"• **User ID:** `{user_id}`\n"
            f"• **Chat ID:** `{chat_id}`\n"
            f"• В базе данных о вас пока нет записи."
        )
    await bot.send_message(
        chat_id=chat_id,
        text=text_out,
        format="markdown",
        attachments=[create_main_menu(user_id)]
    )
    user_states.pop(user_id, None)


# ---------- ОБРАБОТЧИКИ ----------

@dp.bot_started()
async def handle_started(event: BotStarted):
    user = event.user
    user_id = user.user_id
    chat_id = event.chat_id
    name = getattr(user, 'first_name', 'друг')
    username = getattr(user, 'username', None)
    ensure_user(user_id, chat_id, username, name)
    await bot.send_message(
        chat_id=chat_id,
        text=f"👋 Привет, {name}!\n\nВыберите действие:",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_created(Command('start'))
async def cmd_start(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    name = getattr(user, 'first_name', 'друг')
    username = getattr(user, 'username', None)
    ensure_user(user_id, chat_id, username, name)
    user_states.pop(user_id, None)
    user_clarify_data.pop(user_id, None)
    admin_debt_data.pop(user_id, None)
    await bot.send_message(
        chat_id=chat_id,
        text=f"👋 Привет, {name}!\n\nВыберите действие:",
        attachments=[create_main_menu(user_id)]
    )


async def send_help(chat_id: int, user_id: int):
    text = "📋 **Справка**\n\n"
    text += "Используйте кнопки под сообщениями:\n"
    text += "• 📋 Условия – условия по регионам\n"
    if user_id == ADMIN_ID:
        text += "• 💳 Долг – проверить задолженность по любому ФИО\n"
    else:
        text += "• 💳 Долг – проверить задолженность\n"
    text += "• ℹ️ О боте – информация о боте\n"
    text += "• 💬 Поддержка – связаться с администратором\n"
    text += "• ❓ Справка – эта справка\n"
    if user_id == ADMIN_ID:
        text += "• 📢 Новости – сделать рассылку (только для админа)\n"
        text += "• 👤 Данные контакта – получить данные пользователя (только для админа)\n"

    text += "\n**Доступные команды:**\n"
    text += "• /start — Главное меню\n"
    text += "• /help — Эта справка\n"
    text += "• /debt — Узнать долг\n"
    text += "• /change_name — Сменить ФИО\n"
    text += "• /cancel — Отменить текущее действие\n"
    if user_id == ADMIN_ID:
        text += "• /broadcast — Начать рассылку (только админ)\n"
        text += "• /reply — Ответить пользователю (только админ)\n"
        text += "• /userinfo — Данные контакта (только админ)\n"
        text += "• /refresh_debt — Проверить таблицы сверки сейчас (только админ)\n"
        text += "• /debt_status — Статус сверки (только админ)\n"

    await bot.send_message(
        chat_id=chat_id,
        text=text,
        format="markdown",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_created(Command('help'))
async def cmd_help(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await send_help(chat_id, user_id)


@dp.message_callback(F.callback.payload == "help_command")
async def callback_help_command(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="❓ Справка")
    await send_help(chat_id, user_id)


@dp.message_created(Command('change_name'))
async def cmd_change_name(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await start_change_name(user_id, chat_id, event)


async def start_change_name(user_id: int, chat_id: int, event):
    clear_user_full_name(user_id)
    user_states.pop(user_id, None)
    user_clarify_data.pop(user_id, None)
    admin_debt_data.pop(user_id, None)
    user_states[user_id] = UserState.DEBT_NAME_MODE
    await bot.send_message(
        chat_id=chat_id,
        text="🔄 **Смена ФИО**\n\n"
             "Введите ваши ФИО, на которые оформлена основная доверенность.\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )
    if hasattr(event, 'answer'):
        await event.answer(notification="🔄 Введите новые ФИО")


@dp.message_created(Command('debt'))
async def cmd_debt(event: MessageCreated):
    await handle_debt_request(event.message.sender.user_id, event.message.recipient.chat_id, event)


@dp.message_callback(F.callback.payload == "debt")
async def callback_debt(event: MessageCallback):
    user_id = event.callback.user.user_id
    chat_id = event.message.recipient.chat_id
    await handle_debt_request(user_id, chat_id, event)


async def handle_debt_request(user_id: int, chat_id: int, event):
    ensure_user(user_id, chat_id, None, None)

    if user_id == ADMIN_ID:
        user_states[user_id] = UserState.ADMIN_DEBT_MODE
        user_clarify_data.pop(user_id, None)
        admin_debt_data.pop(user_id, None)
        await bot.send_message(
            chat_id=chat_id,
            text="💳 **Проверка долга клиента**\n\n"
                 "Введите ФИО контакта, для которого нужно узнать задолженность.\n"
                 "Например: *Иванов Иван* или *Иванов Иван Иванович*\n\n"
                 "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
            format="markdown",
            attachments=[create_cancel_menu()]
        )
        if hasattr(event, 'answer'):
            await event.answer(notification="💳 Введите ФИО клиента")
        return

    users = load_users()
    user_data = users.get(str(user_id))
    full_name = user_data.get('full_name') if user_data else None

    if full_name:
        await process_debt_search(user_id, chat_id, full_name)
    else:
        user_states[user_id] = UserState.DEBT_NAME_MODE
        user_clarify_data.pop(user_id, None)
        await bot.send_message(
            chat_id=chat_id,
            text="💳 **Проверка долга**\n\n"
                 "Укажите ваши ФИО, на которые оформлена основная доверенность.\n"
                 "Например: *Иванов Иван* или *Иванов Иван Иванович*\n\n"
                 "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
            format="markdown",
            attachments=[create_cancel_menu()]
        )
        if hasattr(event, 'answer'):
            await event.answer(notification="💳 Введите ФИО")


async def process_debt_search(user_id: int, chat_id: int, full_name: str):
    # Поиск с автовосстановлением (скачивает таблицы, если кэша нет)
    matches = await search_debt_with_recovery(full_name)

    if matches is None:
        user_states.pop(user_id, None)
        await bot.send_message(
            chat_id=chat_id,
            text="⚠️ **Сервис проверки долга временно недоступен.**\n\n"
                 "Пожалуйста, попробуйте позже или обратитесь к менеджеру.",
            attachments=[create_main_menu(user_id)]
        )
        return

    if not matches:
        await bot.send_message(
            chat_id=chat_id,
            text="❌ **Доверенность с данными ФИО не найдена.**\n\n"
                 "Пожалуйста, проверьте правильность ввода или обратитесь к менеджеру.\n"
                 "Вы можете сменить ФИО с помощью кнопки ниже.",
            format="markdown",
            attachments=[create_debt_result_menu()]
        )
        return

    groups = group_matches(matches)

    if len(groups) == 1:
        _, items = next(iter(groups.items()))
        display = clean_name(items[0]['name'])
        set_user_full_name(user_id, display)
        await send_debt_result(chat_id, user_id, matches)
        return

    base = ' '.join(_norm_word(w) for w in full_name.strip().split()[:2])
    variants = {key: clean_name(items[0]['name']) for key, items in groups.items()}

    user_clarify_data[user_id] = {'base': base, 'variants': variants}
    user_states[user_id] = UserState.DEBT_CLARIFY_MODE

    variants_text = "\n".join([f"• {name}" for name in variants.values()])
    await bot.send_message(
        chat_id=chat_id,
        text=f"🔍 **Найдено несколько вариантов:**\n\n{variants_text}\n\n"
             "Уточните, пожалуйста, ваше **отчество** (или введите полное ФИО).\n"
             "Например: *Иванович* или *Иванов Иван Иванович*\n\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


async def process_admin_debt_search(user_id: int, chat_id: int, full_name: str):
    # Поиск с автовосстановлением
    matches = await search_debt_with_recovery(full_name)

    if matches is None:
        await bot.send_message(
            chat_id=chat_id,
            text="⚠️ **Сервис проверки долга временно недоступен.**\n\nПопробуйте позже.",
            attachments=[create_main_menu(user_id)]
        )
        return

    if not matches:
        await bot.send_message(
            chat_id=chat_id,
            text=f"❌ **Доверенность с ФИО «{full_name}» не найдена.**\n\n"
                 "Проверьте ввод и попробуйте ещё раз.",
            format="markdown",
            attachments=[create_admin_debt_menu()]
        )
        return

    groups = group_matches(matches)

    if len(groups) == 1:
        await send_debt_result(chat_id, user_id, matches, menu=create_admin_debt_menu())
        return

    base = ' '.join(_norm_word(w) for w in full_name.strip().split()[:2])
    variants = {key: clean_name(items[0]['name']) for key, items in groups.items()}
    admin_debt_data[user_id] = {'base': base, 'variants': variants}
    user_states[user_id] = UserState.ADMIN_DEBT_CLARIFY_MODE

    variants_text = "\n".join([f"• {name}" for name in variants.values()])
    await bot.send_message(
        chat_id=chat_id,
        text=f"🔍 **Найдено несколько вариантов:**\n\n{variants_text}\n\n"
             "Уточните **отчество** клиента (или введите полное ФИО).\n"
             "Например: *Иванович* или *Иванов Иван Иванович*\n\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


@dp.message_callback(F.callback.payload == "change_name")
async def callback_change_name(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await start_change_name(user_id, chat_id, event)


@dp.message_callback(F.callback.payload == "admin_debt_again")
async def callback_admin_debt_again(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id

    if user_id != ADMIN_ID:
        await event.answer(notification="⛔ Доступно только администратору")
        return

    await handle_debt_request(user_id, chat_id, event)


# ---------- КОМАНДА /USERINFO (ТОЛЬКО АДМИН) ----------

@dp.message_created(Command('userinfo'))
async def cmd_userinfo(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))

    if user_id != ADMIN_ID:
        await event.message.answer("⛔ Эта команда доступна только администратору.")
        return

    user_states[user_id] = UserState.WAITING_FOR_FORWARD
    await event.message.answer(
        "👤 **Получение информации о контакте**\n\n"
        "• Перешлите мне сообщение от другого контакта — покажу его данные.\n"
        "• Напишите любое сообщение — покажу ваши данные.\n\n"
        "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


@dp.message_callback(F.callback.payload == "userinfo")
async def callback_userinfo(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))

    if user_id != ADMIN_ID:
        await event.answer(notification="⛔ Доступно только администратору")
        await bot.send_message(
            chat_id=chat_id,
            text="⛔ Эта функция доступна только администратору.",
            attachments=[create_main_menu(user_id)]
        )
        return

    user_states[user_id] = UserState.WAITING_FOR_FORWARD
    await event.answer(notification="👤 Ожидаю ввод")
    await bot.send_message(
        chat_id=chat_id,
        text="👤 **Получение информации о контакте**\n\n"
             "• Перешлите мне сообщение от другого контакта — покажу его данные.\n"
             "• Напишите любое сообщение — покажу ваши данные.\n\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


# ---------- АДМИН-КОМАНДЫ СВЕРКИ ----------

@dp.message_created(Command('refresh_debt'))
async def cmd_refresh_debt(event: MessageCreated):
    """Ручная проверка в любой день. Триггеры проверяются как в понедельник:
    если оба обновились — зачёт + рассылка."""
    user_id = event.message.sender.user_id
    chat_id = event.message.recipient.chat_id

    if user_id != ADMIN_ID:
        await event.message.answer("⛔ Эта команда только для администратора.")
        return

    if not DEBT_SHEET_IDS:
        await event.message.answer("⚠️ DEBT_SHEET_IDS не задан в настройках.")
        return

    await event.message.answer("🔎 Проверяю таблицы…")
    status = await scheduled_check(force_weekday=True)
    await event.message.answer(status)


@dp.message_created(Command('debt_status'))
async def cmd_debt_status(event: MessageCreated):
    user_id = event.message.sender.user_id
    chat_id = event.message.recipient.chat_id

    if user_id != ADMIN_ID:
        await event.message.answer("⛔ Эта команда только для администратора.")
        return

    if not DEBT_SHEET_IDS:
        await event.message.answer("❌ DEBT_SHEET_IDS не задан.")
        return

    now = datetime.now(MSK)
    wd_names = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]

    lines = ["📊 **Статус сверки**\n"]
    lines.append(f"• Сейчас: {wd_names[now.weekday()]}, {now:%H:%M} МСК")
    if now.weekday() in (5, 6):
        lines.append("• Сегодня выходной — плановых проверок нет")
    elif now.weekday() == 0:
        lines.append(f"• Расписание: триггеры + сверка каждые {DEBT_CHECK_MONDAY_MINUTES} мин")
    else:
        lines.append(f"• Расписание: сверка каждые {DEBT_CHECK_WEEKDAY_MINUTES} мин")
    lines.append(f"• Лист с долгами: «{DEBT_DATA_SHEET}»")
    lines.append(f"• Лист-триггер: «{DEBT_SOURCE_SHEET}»")

    st = _debt_state
    today = now.date().isoformat()
    if st['trig_committed_on'] == today:
        lines.append("• Триггеры: ✅ недельное обновление сегодня зачтено — до следующего пн не проверяются\n")
    else:
        lines.append("• Триггеры: ожидание обновления (проверяются только в пн)\n")

    for i, sid in enumerate(DEBT_SHEET_IDS, 1):
        parts = [f"**Таблица {i}**"]
        parts.append(f"• ID: `{sid[:12]}…`")
        tb = st['trig_baseline']
        lt = st['last_sigs_t']
        if tb and lt and i - 1 < len(lt) and lt[i - 1]:
            if tb[i - 1] != lt[i - 1]:
                parts.append("• Триггер: 🟡 изменился (ждём остальные)")
            else:
                parts.append("• Триггер: ✅ без изменений")
        db = st['data_baseline']
        ld = st['last_sigs_d']
        if db and ld and i - 1 < len(ld) and ld[i - 1]:
            if db[i - 1] != ld[i - 1]:
                parts.append("• Сверка: 🟡 в Google новее активной (будет подхвачена)")
            else:
                parts.append("• Сверка: ✅ актуальна")
        lines.append("\n".join(parts))

    if st['last_check']:
        lines.append(f"\n• Последняя проверка: {st['last_check']:%d.%m.%Y %H:%M:%S} МСК")

    rows = _load_debt_rows()
    loaded_at = _debt_cache.get('loaded_at')
    lines.append(f"• Записей в активной сверке: {len(rows) if rows is not None else '⚠️ файлы не читаются'}")
    if loaded_at:
        lines.append(f"• Сверка прочитана ботом: {loaded_at:%d.%m %H:%M:%S}")

    await event.message.answer("\n".join(lines), format="markdown")


# ---------- ОСТАЛЬНЫЕ ОБРАБОТЧИКИ ----------

@dp.message_callback(F.callback.payload == "info")
async def callback_info(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="ℹ️ О боте")
    await bot.send_message(
        chat_id=chat_id,
        text="ℹ️ **О боте**\n\nВерсия: 1.3\nНазначение: автоматизация новостных рассылок для подписчиков и предоставление общей информации по сотрудничеству.\nРазработчик: Александр К.",
        format="markdown",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_callback(F.callback.payload == "conditions")
async def callback_conditions(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="📋 Выбор региона")
    await bot.send_message(
        chat_id=chat_id,
        text="📋 **Выберите ваш регион**\n\nДля просмотра условий выберите нужный регион:",
        format="markdown",
        attachments=[create_regions_menu()]
    )


@dp.message_callback(F.callback.payload == "region_1")
async def callback_region_1(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="📍 Краснодар")
    text = (
        "📋 **Условия по Краснодару**\n\n"
        "1. Товар строго для кормления животных.\n"
        "2. Требуется ежедневный вывоз, либо по согласованному графику с директором магазина.\n"
        "3. Перед приездом нужно оповестить директора магазина.\n"
        "4. Гниль и плесень не забираем. Перед взвешиванием можно отложить гниль в сторону, определив по внешнему виду.\n"
        "5. Тара под списание должна быть своей (тара из магазина возвратная).\n"
        "6. Оплата производится по накладным (с понедельника по воскресенье включительно) в воскресенье/понедельник на р/счёт или по qr-коду.\n"
        "7. При жалобах с магазина, задержках оплаты и других нарушениях условий доверенность аннулируется.\n"
        "8. По вопросам или по сложным ситуациям можно сообщить по номеру: +7 953 641-69-93 Александр.\n\n"
        "Для оформления доверенности необходимы полные фамилия, имя, отчество водителя, серия и номер паспорта, кем и когда выдан, а также заявление, подтверждающее наличие животных.\n"
        "Заявление заполняется каждым водителем, который будет забирать списание. Заявление можно написать от руки на А4 по образцу, поставить подпись, и фото заполненного заявления прислать мне.\n"
        "Доверенность оформляется на человека, который будет забирать просрок.\n\n"
        "• Цена: 15 руб. за 1 кг продукции;\n"
        "• Ассортимент: Хлеб, хлебобулочная продукция, овощи, фрукты, молочная, мясная гастрономия и бакалея;\n"
    )
    await bot.send_message(chat_id=chat_id, text=text, format="markdown")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите нужный документ:**",
        format="markdown",
        attachments=[create_materials_menu("region1")]
    )


@dp.message_callback(F.callback.payload == "region_2")
async def callback_region_2(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="📍 КМВ, КЧР")
    text = (
        "📋 **Условия по Ставропольскому краю и КЧР**\n\n"
        "1. Товар строго для кормления животных.\n"
        "2. Требуется ежедневный вывоз, либо по согласованному графику с директором магазина.\n"
        "3. Перед приездом нужно оповестить директора магазина.\n"
        "4. Гниль и плесень не забираем. Перед взвешиванием можно отложить гниль в сторону, определив по внешнему виду.\n"
        "5. Тара под списание должна быть своей (тара из магазина возвратная).\n"
        "6. Оплата производится по накладным (с понедельника по воскресенье включительно) в воскресенье/понедельник на р/счёт или по qr-коду.\n"
        "7. При жалобах с магазина, задержках оплаты и других нарушениях условий доверенность аннулируется.\n"
        "8. По вопросам или по сложным ситуациям можно сообщить по номеру: +7 953 641-69-93 Александр.\n\n"
        "Для оформления доверенности необходимы полные фамилия, имя, отчество водителя, серия и номер паспорта, кем и когда выдан, а также заявление, подтверждающее наличие животных.\n"
        "Заявление заполняется каждым водителем, который будет забирать списание. Заявление можно написать от руки на А4 по образцу, поставить подпись, и фото заполненного заявления прислать мне.\n"
        "Доверенность оформляется на человека, который будет забирать просрок.\n\n"
        "• Цена: 17 руб. за 1 кг продукции;\n"
        "• Ассортимент: Хлеб, хлебобулочная продукция, овощи, фрукты, молочная, мясная гастрономия и бакалея;\n"
    )
    await bot.send_message(chat_id=chat_id, text=text, format="markdown")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите нужный документ:**",
        format="markdown",
        attachments=[create_materials_menu("region2")]
    )


@dp.message_callback(F.callback.payload == "region1_app")
async def callback_region1_app(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="📄 Заявление")
    await send_image(chat_id, REGION1_APP, "📄 Заявление для Краснодара:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region1")]
    )


@dp.message_callback(F.callback.payload == "region1_qr")
async def callback_region1_qr(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="📱 QR-код")
    await send_image(chat_id, REGION1_QR, "📱 QR-код для Краснодара:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region1")]
    )


@dp.message_callback(F.callback.payload == "region1_bank")
async def callback_region1_bank(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="💳 Реквизиты")
    await send_image(chat_id, REGION1_BANK, "💳 Реквизиты для Краснодара:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region1")]
    )


@dp.message_callback(F.callback.payload == "region2_app")
async def callback_region2_app(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="📄 Заявление")
    await send_image(chat_id, REGION2_APP, "📄 Заявление для КМВ, КЧР:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region2")]
    )


@dp.message_callback(F.callback.payload == "region2_qr")
async def callback_region2_qr(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="📱 QR-код")
    await send_image(chat_id, REGION2_QR, "📱 QR-код для КМВ, КЧР:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region2")]
    )


@dp.message_callback(F.callback.payload == "region2_bank")
async def callback_region2_bank(event: MessageCallback):
    chat_id = event.message.recipient.chat_id
    await event.answer(notification="💳 Реквизиты")
    await send_image(chat_id, REGION2_BANK, "💳 Реквизиты для КМВ, КЧР:")
    await bot.send_message(
        chat_id=chat_id,
        text="📎 **Выберите другой документ:**",
        format="markdown",
        attachments=[create_materials_menu("region2")]
    )


@dp.message_callback(F.callback.payload == "back_to_regions")
async def callback_back_to_regions(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="🔙 Назад к регионам")
    await bot.send_message(
        chat_id=chat_id,
        text="📋 **Выберите регион:**",
        format="markdown",
        attachments=[create_regions_menu()]
    )


@dp.message_callback(F.callback.payload == "back_to_main")
async def callback_back_to_main(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    await event.answer(notification="🔙 Возврат")
    await bot.send_message(
        chat_id=chat_id,
        text="Главное меню:",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_callback(F.callback.payload == "support")
async def callback_support(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    user_states[user_id] = UserState.SUPPORT_MODE
    await event.answer(notification="💬 Поддержка активирована")
    await bot.send_message(
        chat_id=chat_id,
        text="💬 **Поддержка**\n\n"
             "Напишите ваше сообщение, и я перешлю его администратору.\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


@dp.message_callback(F.callback.payload == "broadcast")
async def callback_broadcast(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))

    if user_id != ADMIN_ID:
        await event.answer(notification="⛔ Доступно только администратору")
        await bot.send_message(
            chat_id=chat_id,
            text="⛔ Эта функция доступна только администратору.",
            attachments=[create_main_menu(user_id)]
        )
        return

    user_states[user_id] = UserState.BROADCAST_MODE
    await event.answer(notification="📢 Режим рассылки")
    await bot.send_message(
        chat_id=chat_id,
        text="📢 **Рассылка новостей**\n\n"
             "Введите текст, который нужно разослать всем пользователям.\n"
             "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


@dp.message_callback(F.callback.payload == "cancel_action")
async def callback_cancel_action(event: MessageCallback):
    user = event.callback.user
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    user_states.pop(user_id, None)
    user_clarify_data.pop(user_id, None)
    admin_debt_data.pop(user_id, None)
    await event.answer(notification="❌ Действие отменено")
    await bot.send_message(
        chat_id=chat_id,
        text="❌ Вы отменили текущее действие.",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_created(Command('broadcast'))
async def cmd_broadcast(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))

    if user_id != ADMIN_ID:
        await event.message.answer("⛔ Эта команда только для администратора.")
        return

    user_states[user_id] = UserState.BROADCAST_MODE
    await event.message.answer(
        "📢 **Рассылка новостей**\n\n"
        "Введите текст, который нужно разослать всем пользователям.\n"
        "Для отмены нажмите кнопку «Отмена» или введите команду /cancel.",
        format="markdown",
        attachments=[create_cancel_menu()]
    )


@dp.message_created(Command('cancel'))
async def cmd_cancel(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))
    user_states.pop(user_id, None)
    user_clarify_data.pop(user_id, None)
    admin_debt_data.pop(user_id, None)
    await event.message.answer(
        "❌ Действие отменено.",
        attachments=[create_main_menu(user_id)]
    )


@dp.message_created(Command('reply'))
async def cmd_reply(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    ensure_user(user_id, chat_id, getattr(user, 'username', None), getattr(user, 'first_name', None))

    if user_id != ADMIN_ID:
        await event.message.answer("⛔ Эта команда только для администратора.")
        return

    body = getattr(event.message, 'body', None)
    text = (getattr(body, 'text', None) or '') if body else ''

    match = re.match(r'^/reply\s+(\d+)\s+(.+)$', text, re.DOTALL)
    if not match:
        await event.message.answer(
            "❌ Неверный формат.\n"
            "Используйте: `/reply <user_id> <текст>`\n"
            "Пример: `/reply 123456789 Привет, ваш вопрос решён!`",
            format="markdown"
        )
        return

    target_user_id = int(match.group(1))
    reply_text = match.group(2).strip()
    if not reply_text:
        await event.message.answer("❌ Текст ответа не может быть пустым.")
        return

    users = load_users()
    target_data = users.get(str(target_user_id))
    if not target_data:
        await event.message.answer(f"❌ Пользователь с ID {target_user_id} не найден в базе.")
        return

    target_chat_id = target_data.get('chat_id')
    if not target_chat_id:
        await event.message.answer(f"❌ Для пользователя {target_user_id} не найден chat_id.")
        return

    try:
        await bot.send_message(
            chat_id=target_chat_id,
            text=f"📨 **Ответ администратора:**\n\n{reply_text}",
            format="markdown"
        )
        await event.message.answer(f"✅ Ответ отправлен пользователю {target_user_id}.")
        logger.info(f"Администратор ответил пользователю {target_user_id}")
    except Exception as e:
        logger.error(f"Ошибка отправки ответа пользователю {target_user_id}: {e}")
        await event.message.answer("❌ Не удалось отправить ответ. Проверьте ID пользователя или права бота.")


# ---------- ВСПОМОГАТЕЛЬНАЯ ФУНКЦИЯ ДЛЯ КАРТИНОК ----------

async def send_image(chat_id: int, image_path_or_url: str, caption: str = ""):
    """CHANGED: URL — отправка по ссылке как раньше; локальный путь —
    загрузка файла в MAX через upload_media и отправка готового вложения."""
    if not image_path_or_url:
        logger.warning("Путь/URL изображения не указан")
        return
    try:
        if image_path_or_url.startswith("http"):
            # как раньше: вложение по URL
            image_attachment = Attachment(type="image", payload={"url": image_path_or_url})
            await bot.send_message(chat_id=chat_id, text=caption,
                                   attachments=[image_attachment])
        else:
            # локальный файл: читаем и загружаем в MAX
            if not os.path.exists(image_path_or_url):
                logger.warning(f"Файл не найден: {image_path_or_url}")
                return
            with open(image_path_or_url, "rb") as f:
                data = f.read()
            media = InputMediaBuffer(
                buffer=data,
                filename=os.path.basename(image_path_or_url),
                type="image"
            )
            uploaded = await bot.upload_media(media)
            # upload_media может вернуть готовый Attachment или токен-строку —
            # обрабатываем оба варианта
            if isinstance(uploaded, Attachment):
                attachments = [uploaded]
            elif isinstance(uploaded, str):
                attachments = [Attachment(type="image", payload={"token": uploaded})]
            else:
                attachments = [uploaded]
            await bot.send_message(chat_id=chat_id, text=caption, attachments=attachments)

        logger.info(f"Изображение отправлено: {image_path_or_url}")
    except Exception as e:
        logger.error(f"Ошибка отправки изображения {image_path_or_url}: {e}")


# ---------- ОСНОВНОЙ ОБРАБОТЧИК ТЕКСТА ----------
@dp.message_created(F.message.body.text)
async def handle_text(event: MessageCreated):
    text = event.message.body.text
    if not text:
        return

    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    username = getattr(user, 'username', None)
    name = getattr(user, 'first_name', 'Пользователь')
    ensure_user(user_id, chat_id, username, name)

    if text.startswith('/'):
        first_word = text.split(maxsplit=1)[0].split('@')[0].lower()
        if first_word not in KNOWN_COMMANDS:
            await event.message.answer(
                "❓ Неизвестная команда. Список доступных команд: /help",
                attachments=[create_main_menu(user_id)]
            )
        return

    state = user_states.get(user_id)

    # ---------- РЕЖИМ ОЖИДАНИЯ (USERINFO) ----------
    if state == UserState.WAITING_FOR_FORWARD:
        await process_userinfo_input(event, user_id, chat_id, text)
        return

    # ---------- АДМИН: ВВОД ФИО КЛИЕНТА ----------
    if state == UserState.ADMIN_DEBT_MODE:
        full_name = text.strip()
        if len(full_name.split()) < 2:
            await event.message.answer(
                "⚠️ **Нужно ввести хотя бы фамилию и имя.**\n"
                "Например: *Иванов Иван* или *Иванов Иван Иванович*\n"
                "Вы можете нажать «Отмена» для выхода.",
                format="markdown",
                attachments=[create_cancel_menu()]
            )
            return

        user_states.pop(user_id, None)
        await process_admin_debt_search(user_id, chat_id, full_name)
        return

    # ---------- АДМИН: УТОЧНЕНИЕ ОТЧЕСТВА КЛИЕНТА ----------
    if state == UserState.ADMIN_DEBT_CLARIFY_MODE:
        data = admin_debt_data.get(user_id)
        if not data:
            user_states.pop(user_id, None)
            await event.message.answer(
                "❌ **Ошибка: данные для уточнения не найдены.**\n"
                "Начните проверку заново.",
                attachments=[create_main_menu(user_id)]
            )
            return

        lw = [_norm_word(w) for w in text.strip().split()]
        base_words = str(data['base']).split()
        variants: Dict[str, str] = data['variants']

        candidates: List[str] = []
        if len(lw) == 1:
            candidates.append(' '.join(base_words + lw))
        else:
            candidates.append(' '.join(lw[:3]))
            candidates.append(' '.join(lw[:2]))

        found_key = next((c for c in candidates if c in variants), None)
        if found_key is None:
            entered_norm = ' '.join(lw[:3])
            for key, display in variants.items():
                if ' '.join(_norm_word(w) for w in display.split()[:3]) == entered_norm:
                    found_key = key
                    break

        if found_key:
            matches = await search_debt_with_recovery(found_key)
            user_states.pop(user_id, None)
            admin_debt_data.pop(user_id, None)
            if matches:
                groups = group_matches(matches)
                matches = groups.get(found_key, matches)
                await send_debt_result(chat_id, user_id, matches, menu=create_admin_debt_menu())
            else:
                await bot.send_message(
                    chat_id=chat_id,
                    text="❌ **Не удалось найти долг по уточнённому ФИО.**",
                    attachments=[create_admin_debt_menu()]
                )
        else:
            await event.message.answer(
                "❌ **Вариант не найден.**\n\n"
                "Проверьте ввод и попробуйте снова.\n"
                "Доступные варианты:\n" + "\n".join([f"• {n}" for n in variants.values()]),
                format="markdown",
                attachments=[create_cancel_menu()]
            )
        return

    # ---------- РЕЖИМ УТОЧНЕНИЯ ОТЧЕСТВА (пользователь) ----------
    if state == UserState.DEBT_CLARIFY_MODE:
        entered = text.strip()
        data = user_clarify_data.get(user_id)
        if not data:
            user_states.pop(user_id, None)
            await event.message.answer(
                "❌ **Ошибка: данные для уточнения не найдены.**\n"
                "Пожалуйста, начните проверку долга заново.",
                attachments=[create_main_menu(user_id)]
            )
            return

        lw = [_norm_word(w) for w in entered.split()]
        base_words = str(data['base']).split()
        variants: Dict[str, str] = data['variants']

        candidates: List[str] = []
        if len(lw) == 1:
            candidates.append(' '.join(base_words + lw))
        else:
            candidates.append(' '.join(lw[:3]))
            candidates.append(' '.join(lw[:2]))

        found_key = next((c for c in candidates if c in variants), None)
        if found_key is None:
            entered_norm = ' '.join(lw[:3])
            for key, display in variants.items():
                if ' '.join(_norm_word(w) for w in display.split()[:3]) == entered_norm:
                    found_key = key
                    break

        if found_key:
            display_name = variants[found_key]
            set_user_full_name(user_id, display_name)
            matches = await search_debt_with_recovery(found_key)
            user_states.pop(user_id, None)
            user_clarify_data.pop(user_id, None)
            if matches:
                groups = group_matches(matches)
                matches = groups.get(found_key, matches)
                await send_debt_result(chat_id, user_id, matches)
            else:
                await bot.send_message(
                    chat_id=chat_id,
                    text="❌ **Не удалось найти долг по уточнённому ФИО.**\n"
                         "Пожалуйста, обратитесь к администратору.",
                    attachments=[create_main_menu(user_id)]
                )
        else:
            await event.message.answer(
                "❌ **Вариант не найден.**\n\n"
                "Пожалуйста, проверьте правильность ввода и попробуйте снова.\n"
                "Доступные варианты:\n" + "\n".join([f"• {n}" for n in variants.values()]),
                format="markdown",
                attachments=[create_cancel_menu()]
            )
        return

    # ---------- РЕЖИМ ВВОДА ФИО (пользователь) ----------
    if state == UserState.DEBT_NAME_MODE:
        full_name = text.strip()
        if len(full_name.split()) < 2:
            await event.message.answer(
                "⚠️ **Нужно ввести хотя бы фамилию и имя.**\n"
                "Например: *Иванов Иван* или *Иванов Иван Иванович*\n"
                "Вы можете нажать «Отмена» для выхода.",
                format="markdown",
                attachments=[create_cancel_menu()]
            )
            return

        set_user_full_name(user_id, full_name)
        user_states.pop(user_id, None)

        await process_debt_search(user_id, chat_id, full_name)
        return

    # ---------- РЕЖИМ ПОДДЕРЖКИ ----------
    if state == UserState.SUPPORT_MODE:
        admin_chat_id = None
        if ADMIN_ID:
            users = load_users()
            admin_data = users.get(str(ADMIN_ID))
            admin_chat_id = admin_data.get('chat_id') if admin_data else None

        if not admin_chat_id:
            user_states.pop(user_id, None)
            logger.error("Поддержка недоступна: администратор не настроен")
            await event.message.answer(
                "⚠️ К сожалению, поддержка сейчас недоступна. Попробуйте позже.",
                attachments=[create_main_menu(user_id)]
            )
            return

        username_str = f"@{username}" if username else "без username"
        try:
            await bot.send_message(
                chat_id=admin_chat_id,
                text=f"📩 **Новое обращение в поддержку**\n\n"
                     f"От: {name} ({username_str})\n"
                     f"ID пользователя: `{user_id}`\n"
                     f"Сообщение:\n{text}",
                format="markdown"
            )
        except Exception as e:
            logger.error(f"Не удалось переслать сообщение админу: {e}")
            await event.message.answer(
                "⚠️ Ошибка отправки. Попробуйте позже.",
                attachments=[create_cancel_menu()]
            )
            return

        user_states.pop(user_id, None)
        logger.info(f"Обращение в поддержку от user_id={user_id} переслано администратору")
        await event.message.answer(
            "✅ Сообщение отправлено!\n"
            "Мы ответим в ближайшее время.",
            attachments=[create_main_menu(user_id)]
        )
        return

    # ---------- РЕЖИМ РАССЫЛКИ ----------
    if state == UserState.BROADCAST_MODE:
        if user_id != ADMIN_ID:
            user_states.pop(user_id, None)
            await event.message.answer(
                "⛔ Доступно только администратору.",
                attachments=[create_main_menu(user_id)]
            )
            return

        users = load_users()
        sent_count = 0
        error_count = 0
        logger.info(f"Начинаем рассылку от админа {user_id}, пользователей: {len(users)}")

        for uid, user_data in users.items():
            chat_id_user = user_data.get('chat_id')
            if not chat_id_user:
                continue
            try:
                await bot.send_message(
                    chat_id=chat_id_user,
                    text=f"📢 **Новость от администратора:**\n\n{text}",
                    format="markdown"
                )
                sent_count += 1
            except Exception as e:
                logger.error(f"Не удалось отправить новость пользователю {uid}: {e}")
                error_count += 1
            await asyncio.sleep(0.2)

        user_states.pop(user_id, None)
        await event.message.answer(
            f"✅ **Рассылка завершена.**\n"
            f"Отправлено: {sent_count}\n"
            f"Ошибок: {error_count}",
            format="markdown",
            attachments=[create_main_menu(user_id)]
        )
        return

    # ---------- АДМИНИСТРАТОР БЕЗ АКТИВНОГО РЕЖИМА ----------
    if user_id == ADMIN_ID:
        await event.message.answer(
            "👋 Вы администратор. Используйте:\n"
            "• /broadcast – для рассылки новостей\n"
            "• /reply <id> <текст> – для ответа пользователю\n"
            "• /userinfo – для получения данных пользователя\n"
            "• /debt или кнопка «Долг» – проверить задолженность по ФИО клиента\n"
            "• /refresh_debt – проверить таблицы сверки сейчас\n"
            "• /debt_status – статус сверки",
            format="markdown",
            attachments=[create_main_menu(user_id)]
        )
        return

    # ---------- ОБЫЧНЫЙ ПОЛЬЗОВАТЕЛЬ ----------
    await event.message.answer(
        "Не понял вас. Используйте кнопки под сообщениями или /help для справки.",
        attachments=[create_main_menu(user_id)]
    )


# ---------- ОБРАБОТЧИК СООБЩЕНИЙ БЕЗ ТЕКСТА ----------
@dp.message_created()
async def handle_non_text(event: MessageCreated):
    user = event.message.sender
    user_id = user.user_id
    chat_id = event.message.recipient.chat_id
    username = getattr(user, 'username', None)
    name = getattr(user, 'first_name', 'Пользователь')

    body = getattr(event.message, 'body', None)
    text = getattr(body, 'text', '') if body else ''

    if text:
        return

    ensure_user(user_id, chat_id, username, name)
    state = user_states.get(user_id)

    if state == UserState.WAITING_FOR_FORWARD:
        await process_userinfo_input(event, user_id, chat_id, text or "")
        return

    if state == UserState.SUPPORT_MODE:
        await event.message.answer(
            "⚠️ Пожалуйста, отправьте сообщение текстом — вложения не пересылаются.",
            attachments=[create_cancel_menu()]
        )
        return

    if state in (UserState.ADMIN_DEBT_MODE, UserState.ADMIN_DEBT_CLARIFY_MODE):
        await event.message.answer(
            "⚠️ Введите ФИО текстом.",
            attachments=[create_cancel_menu()]
        )
        return

    logger.info(f"Сообщение без текста от user_id={user_id}, state={state} — игнор")


# ==================== ЗАПУСК БОТА ====================

async def main():
    logger.info("Запуск бота...")

    await bot.set_commands(
        BotCommand(name="start", description="Главное меню"),
        BotCommand(name="help", description="Справка"),
        BotCommand(name="debt", description="Долг"),
        BotCommand(name="change_name", description="Сменить ФИО"),
        BotCommand(name="userinfo", description="Данные контакта (админ)"),
        BotCommand(name="broadcast", description="Рассылка (только админ)"),
        BotCommand(name="reply", description="Ответ пользователю (только админ)"),
        BotCommand(name="refresh_debt", description="Проверить сверку сейчас (админ)"),
        BotCommand(name="debt_status", description="Статус сверки (админ)"),
        BotCommand(name="cancel", description="Отменить текущее действие"),
    )

    await bot.delete_webhook()

    # Фоновый цикл: стартовая загрузка + расписание пн/вт-пт/сб-вс
    asyncio.create_task(debt_refresh_loop())

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
